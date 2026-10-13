import openpyxl

from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render

from .models import SMSTemplate, SMSCampaign
from .report import generate_sms_campaign_report,generate_batch_sms_fallback_report
from .tasks import process_sms_campaign


# =========================================================
# UPLOAD EXCEL
# =========================================================

def upload_excel(request):

    templates = SMSTemplate.objects.all()

    if request.method != "POST":
        return render(
            request,
            "upload.html",
            {
                "templates": templates,
            },
        )

    uploaded_file = request.FILES.get("file")
    template_id = request.POST.get("template")

    # -----------------------------------------------------
    # FILE VALIDATION
    # -----------------------------------------------------

    if not uploaded_file:
        return render(
            request,
            "upload.html",
            {
                "templates": templates,
                "error": "Please select an Excel file.",
            },
        )

    # -----------------------------------------------------
    # TEMPLATE VALIDATION
    # -----------------------------------------------------

    if not template_id:
        return render(
            request,
            "upload.html",
            {
                "templates": templates,
                "error": "Please select a template.",
            },
        )

    try:

        template = SMSTemplate.objects.get(
            id=template_id
        )

    except SMSTemplate.DoesNotExist:

        return render(
            request,
            "upload.html",
            {
                "templates": templates,
                "error": "Selected template does not exist.",
            },
        )

    try:

        # =================================================
        # TEMPLATE VARIABLES
        # =================================================

        variables = [
            variable.strip().lower()
            for variable in template.variables.split(",")
            if variable.strip()
        ]

        placeholder_count = (
            template.content.count("*")
        )

        if placeholder_count != len(variables):

            return render(
                request,
                "upload.html",
                {
                    "templates": templates,
                    "error": (
                        f"Template '{template.name}' "
                        f"has {placeholder_count} '*' "
                        f"placeholders but "
                        f"{len(variables)} variables."
                    ),
                },
            )

        # =================================================
        # EXCEL VALIDATION
        # =================================================

        workbook = openpyxl.load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )

        try:

            sheet = workbook.active

            first_row = next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                None,
            )

        finally:

            workbook.close()

        # -------------------------------------------------
        # EMPTY EXCEL
        # -------------------------------------------------

        if not first_row:

            return render(
                request,
                "upload.html",
                {
                    "templates": templates,
                    "error": "Excel file is empty.",
                },
            )

        # =================================================
        # NORMALIZE HEADERS
        # =================================================

        headers = [
            str(header).strip().lower()
            if header is not None
            else ""
            for header in first_row
        ]

        # =================================================
        # REQUIRED COLUMNS
        # =================================================

        required_columns = [
            "mobile",
            *variables,
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing_columns:

            return render(
                request,
                "upload.html",
                {
                    "templates": templates,
                    "error": (
                        "Missing Excel columns: "
                        + ", ".join(missing_columns)
                    ),
                },
            )

        # =================================================
        # CREATE CAMPAIGN
        # =================================================

        campaign = SMSCampaign.objects.create(
            name=f"{template.name} Campaign",

            template=template,

            excel_file=uploaded_file,

            status="PENDING",

            total=0,

            processed=0,

            sent=0,

            failed=0,

            skipped=0,

            error_message=None,
        )

        # =================================================
        # START CELERY
        # =================================================

        process_sms_campaign.delay(
            campaign.id
        )

        # =================================================
        # SUCCESS PAGE
        # =================================================

        return render(
            request,
            "success.html",
            {
                "campaign": campaign,
            },
        )

    except Exception as exc:

        return render(
            request,
            "upload.html",
            {
                "templates": templates,
                "error": str(exc),
            },
        )


# =========================================================
# CAMPAIGN STATUS API
# =========================================================

def campaign_status_api(
    request,
    campaign_id,
):
    """
    Return campaign status and counters
    for AJAX polling.
    """

    try:

        campaign = SMSCampaign.objects.get(
            id=campaign_id
        )

    except SMSCampaign.DoesNotExist:

        return JsonResponse(
            {
                "error": "Campaign not found",
            },
            status=404,
        )

    # -----------------------------------------------------
    # COUNTS
    # -----------------------------------------------------

    total = campaign.total or 0

    sent = campaign.sent or 0

    failed = campaign.failed or 0

    skipped = campaign.skipped or 0

    processed = campaign.processed or 0

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    return JsonResponse(
        {
            "id": campaign.id,

            "status": campaign.status,

            "total": total,

            "sent": sent,

            "failed": failed,

            "skipped": skipped,

            "processed": processed,

            "is_completed": (
                campaign.status
                in ["COMPLETED", "FAILED"]
            ),

            "is_processing": (
                campaign.status == "PROCESSING"
            ),

            "is_pending": (
                campaign.status == "PENDING"
            ),

            "error_message": (
                campaign.error_message
            ),
        }
    )


# =========================================================
# ADD SMS TEMPLATE
# =========================================================
from batch_app.models import NotificationType   #s_ms
def ad_template(request):
    """
    Serve SMS template admin page.
    """

    notification_types = NotificationType.objects.filter(
        is_active=True
    ).order_by("name")

    return render(
        request,
        "add_template.html",
        {
            "notification_types": notification_types,
        },
    )


# =========================================================
# DOWNLOAD CAMPAIGN EXCEL REPORT
# =========================================================

def download_sms_campaign_report(
    request,
    campaign_id,
):
    """
    Download one Excel file containing:

    1. Summary
    2. Success
    3. Failed
    4. Skipped
    """

    campaign = get_object_or_404(
        SMSCampaign,
        id=campaign_id,
    )

    # -----------------------------------------------------
    # GENERATE REPORT
    # -----------------------------------------------------

    excel_file = (
        generate_sms_campaign_report(
            campaign.id
        )
    )

    # -----------------------------------------------------
    # FILE NAME
    # -----------------------------------------------------

    filename = (
        f"sms_campaign_"
        f"{campaign.id}_"
        f"report.xlsx"
    )

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    response = HttpResponse(
        excel_file.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from batch_app.models import BatchJob


def batch_sms_fallback_report(request, job_id):

    job = get_object_or_404(
        BatchJob,
        job_id=job_id
    )

    output = generate_batch_sms_fallback_report(
        job.job_id
    )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        f'attachment; filename="'
        f'batch_{job.job_id}_sms_fallback_report.xlsx"'
    )

    return response


def add_notification(request):
    return render(request, 'add_notification.html')
