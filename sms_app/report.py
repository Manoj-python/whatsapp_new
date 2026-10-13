from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import SMSCampaign, SMSLog


def excel_datetime(value):
    if value is None:
        return ""

    if timezone.is_aware(value):
        value = timezone.localtime(value)

    return value.replace(tzinfo=None)


def generate_sms_campaign_report(campaign_id):

    campaign = (
        SMSCampaign.objects
        .select_related("template")
        .get(id=campaign_id)
    )

    workbook = Workbook()

    # =========================================================
    # SUMMARY
    # =========================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    # =========================================================
    # DETAIL SHEETS
    # =========================================================

    success_sheet = workbook.create_sheet("Success")
    failed_sheet = workbook.create_sheet("Failed")
    skipped_sheet = workbook.create_sheet("Skipped")

    # =========================================================
    # GET COUNTS
    # =========================================================

    success_count = SMSLog.objects.filter(
        campaign_id=campaign_id,
        status="sent"
    ).count()

    failed_count = SMSLog.objects.filter(
        campaign_id=campaign_id,
        status="failed"
    ).count()

    skipped_count = SMSLog.objects.filter(
        campaign_id=campaign_id,
        status="skipped"
    ).count()

    # =========================================================
    # SUMMARY
    # =========================================================

    summary_data = [
        ["SMS Campaign Report", ""],
        ["Campaign ID", campaign.id],
        ["Campaign Name", campaign.name or ""],
        ["Template", campaign.template.name],
        ["Template ID", campaign.template.template_id],
        ["Status", campaign.status],
        ["Total", campaign.total],
        ["Processed", campaign.processed],
        ["Success", success_count],
        ["Failed", failed_count],
        ["Skipped", skipped_count],
        ["Created At", excel_datetime(campaign.created_at)],
        ["Started At", excel_datetime(campaign.started_at)],
        ["Completed At", excel_datetime(campaign.completed_at)],
    ]

    for row in summary_data:
        summary_sheet.append(row)

    summary_sheet["A1"].font = Font(
        bold=True,
        size=14
    )

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    summary_sheet.merge_cells("A1:B1")

    for row in summary_sheet.iter_rows():
        if row[0].value != "SMS Campaign Report":
            row[0].font = Font(bold=True)

    summary_sheet.column_dimensions["A"].width = 25
    summary_sheet.column_dimensions["B"].width = 40

    # =========================================================
    # DETAIL SHEET HEADERS
    # =========================================================

    headers = [
        "Mobile",
        "Message",
        "Response",
        "Created At",
    ]

    for sheet in [
        success_sheet,
        failed_sheet,
        skipped_sheet,
    ]:

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 60
        sheet.column_dimensions["C"].width = 60
        sheet.column_dimensions["D"].width = 22

        sheet.freeze_panes = "A2"

    # =========================================================
    # GET ALL LOGS
    # =========================================================

    logs = (
        SMSLog.objects
        .filter(campaign_id=campaign_id)
        .order_by("id")
        .iterator(chunk_size=2000)
    )

    # =========================================================
    # WRITE LOGS TO RESPECTIVE SHEETS
    # =========================================================

    for log in logs:

        row = [
            log.mobile,
            log.message,
            log.response or "",
            excel_datetime(log.created_at),
        ]

        if log.status == "sent":

            success_sheet.append(row)

        elif log.status == "failed":

            failed_sheet.append(row)

        elif log.status == "skipped":

            skipped_sheet.append(row)

    # =========================================================
    # SAVE
    # =========================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


def generate_batch_sms_fallback_report(job_id):

    from batch_app.models import BatchJob

    job = (
        BatchJob.objects
        .select_related("notification_type")
        .get(job_id=job_id)
    )

    workbook = Workbook()

    # =========================================================
    # SUMMARY
    # =========================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    # =========================================================
    # DETAIL SHEETS
    # =========================================================

    success_sheet = workbook.create_sheet("Success")
    failed_sheet = workbook.create_sheet("Failed")

    # =========================================================
    # GET LOGS
    # =========================================================

    logs = (
        SMSLog.objects
        .filter(
            batch_job_id=job.id,
            notification_type_id=job.notification_type_id,
        )
        .order_by("id")
        .iterator(chunk_size=2000)
    )

    # =========================================================
    # COUNTS
    # =========================================================

    success_count = SMSLog.objects.filter(
        batch_job_id=job.id,
        notification_type_id=job.notification_type_id,
        status="sent",
    ).count()

    failed_count = SMSLog.objects.filter(
        batch_job_id=job.id,
        notification_type_id=job.notification_type_id,
        status="failed",
    ).count()

    total_count = success_count + failed_count

    # =========================================================
    # SUMMARY
    # =========================================================

    notification_name = ""

    if job.notification_type:
        notification_name = str(job.notification_type)

    summary_data = [
        ["Batch Job SMS Fallback Report", ""],
        ["Batch Job ID", job.job_id],
        ["Notification Type", notification_name],
        ["Total SMS Fallback", total_count],
        ["SMS Sent", success_count],
        ["SMS Failed", failed_count],
    ]

    for row in summary_data:
        summary_sheet.append(row)

    summary_sheet["A1"].font = Font(
        bold=True,
        size=14
    )

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    summary_sheet.merge_cells("A1:B1")

    for row in summary_sheet.iter_rows():
        if row[0].value != "Batch Job SMS Fallback Report":
            row[0].font = Font(bold=True)

    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 50

    # =========================================================
    # DETAIL HEADERS
    # =========================================================

    headers = [
        "Mobile",
        "Message",
        "Response",
        "Created At",
    ]

    for sheet in [
        success_sheet,
        failed_sheet,
    ]:

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 60
        sheet.column_dimensions["C"].width = 60
        sheet.column_dimensions["D"].width = 22

        sheet.freeze_panes = "A2"

    # =========================================================
    # WRITE LOGS
    # =========================================================

    for log in logs:

        row = [
            log.mobile,
            log.message,
            log.response or "",
            excel_datetime(log.created_at),
        ]

        if log.status == "sent":

            success_sheet.append(row)

        elif log.status == "failed":

            failed_sheet.append(row)

    # =========================================================
    # SAVE
    # =========================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output