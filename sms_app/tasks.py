import openpyxl

from celery import shared_task
from django.utils import timezone

from .models import SMSCampaign, SMSLog
from .utils import (
    create_sms_session,
    send_sms,
)


# ============================================================
# CONFIGURATION
# ============================================================

LOG_BATCH_SIZE = 1000
PROGRESS_BATCH_SIZE = 1000

SKIP_DUPLICATE_MOBILES = True


# ============================================================
# HELPERS
# ============================================================

def normalize_mobile(value):

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    return str(value).strip()


def normalize_header(value):

    if value is None:
        return ""

    return str(value).strip().lower()


def build_message(
    template_content,
    variables,
    variable_indexes,
    row,
):

    message = template_content

    for variable in variables:

        column_index = variable_indexes[variable]

        value = (
            row[column_index]
            if column_index < len(row)
            else None
        )

        if value is None:
            value = ""

        message = message.replace(
            "*",
            str(value),
            1,
        )

    return message


def flush_logs(logs):

    if not logs:
        return []

    SMSLog.objects.bulk_create(
        logs,
        batch_size=LOG_BATCH_SIZE,
    )

    return []


def update_campaign_progress(
    campaign_id,
    processed,
    sent,
    failed,
    skipped,
):

    SMSCampaign.objects.filter(
        id=campaign_id,
    ).update(
        processed=processed,
        sent=sent,
        failed=failed,
        skipped=skipped,
    )


def fail_campaign(
    campaign_id,
    error_message,
):

    SMSCampaign.objects.filter(
        id=campaign_id,
    ).update(
        status="FAILED",
        error_message=str(error_message),
        completed_at=timezone.now(),
    )


# ============================================================
# CLAIM CAMPAIGN
# ============================================================

def claim_campaign(campaign_id):

    updated = (
        SMSCampaign.objects
        .filter(
            id=campaign_id,
            status="PENDING",
        )
        .update(
            status="PROCESSING",
            error_message=None,
            started_at=timezone.now(),
            completed_at=None,
        )
    )

    return updated == 1


# ============================================================
# TEMPLATE
# ============================================================

def get_template_variables(template):

    raw_variables = template.variables or ""

    variables = [
        variable.strip().lower()
        for variable in raw_variables.split(",")
        if variable.strip()
    ]

    if len(variables) != len(set(variables)):

        raise ValueError(
            f"Template '{template.name}' contains "
            f"duplicate variables."
        )

    return variables


def validate_template(
    template,
    variables,
):

    placeholder_count = (
        template.content.count("*")
    )

    if placeholder_count != len(variables):

        raise ValueError(
            f"Template '{template.name}' has "
            f"{placeholder_count} '*' placeholders "
            f"but {len(variables)} variables configured."
        )


# ============================================================
# EXCEL
# ============================================================

def prepare_excel_columns(
    sheet,
    variables,
):

    rows = sheet.iter_rows(
        values_only=True,
    )

    try:

        raw_headers = next(rows)

    except StopIteration:

        raise ValueError(
            "Excel file is empty."
        )

    headers = [
        normalize_header(header)
        for header in raw_headers
    ]

    # ----------------------------------------------
    # Duplicate headers
    # ----------------------------------------------

    non_empty_headers = [
        header
        for header in headers
        if header
    ]

    if len(non_empty_headers) != len(
        set(non_empty_headers)
    ):

        raise ValueError(
            "Excel contains duplicate column names."
        )

    # ----------------------------------------------
    # Mobile
    # ----------------------------------------------

    try:

        mobile_index = headers.index(
            "mobile"
        )

    except ValueError:

        raise ValueError(
            "Excel must contain a 'mobile' column."
        )

    # ----------------------------------------------
    # Variables
    # ----------------------------------------------

    variable_indexes = {}

    for variable in variables:

        try:

            variable_indexes[variable] = (
                headers.index(variable)
            )

        except ValueError:

            raise ValueError(
                f"Excel missing required column: "
                f"{variable}"
            )

    return (
        rows,
        mobile_index,
        variable_indexes,
    )


# ============================================================
# CELERY TASK
# ============================================================

@shared_task(bind=True)
def process_sms_campaign(
    self,
    campaign_id,
):

    # ========================================================
    # CLAIM
    # ========================================================

    if not claim_campaign(campaign_id):

        return {
            "status": "SKIPPED",
            "campaign_id": campaign_id,
            "reason": (
                "Campaign is not in PENDING status."
            ),
        }

    excel_file = None
    workbook = None
    session = None

    logs = []

    processed = 0
    sent = 0
    failed = 0
    skipped = 0

    processed_mobiles = set()

    try:

        # ====================================================
        # LOAD CAMPAIGN
        # ====================================================

        try:

            campaign = (
                SMSCampaign.objects
                .select_related("template")
                .get(id=campaign_id)
            )

        except SMSCampaign.DoesNotExist:

            raise ValueError(
                f"SMS campaign {campaign_id} "
                f"does not exist."
            )

        template = campaign.template

        if template is None:

            raise ValueError(
                "Campaign template is not configured."
            )

        # ====================================================
        # TEMPLATE VALIDATION
        # ====================================================

        variables = get_template_variables(
            template
        )

        validate_template(
            template,
            variables,
        )

        # ====================================================
        # EXCEL
        # ====================================================

        if not campaign.excel_file:

            raise ValueError(
                "Campaign Excel file is not configured."
            )

        excel_file = campaign.excel_file.open(
            "rb"
        )

        workbook = openpyxl.load_workbook(
            excel_file,
            read_only=True,
            data_only=True,
        )

        sheet = workbook.active

        if sheet is None:

            raise ValueError(
                "Excel workbook has no active sheet."
            )

        # ====================================================
        # COLUMNS
        # ====================================================

        (
            rows,
            mobile_index,
            variable_indexes,
        ) = prepare_excel_columns(
            sheet,
            variables,
        )

        # ====================================================
        # TOTAL
        # ====================================================

        total = max(
            (sheet.max_row or 1) - 1,
            0,
        )

        SMSCampaign.objects.filter(
            id=campaign_id,
        ).update(
            total=total,
            processed=0,
            sent=0,
            failed=0,
            skipped=0,
            error_message=None,
        )

        # ====================================================
        # SMS SESSION
        # ====================================================

        session = create_sms_session()

        # ====================================================
        # PROCESS
        # ====================================================

        for row in rows:

            processed += 1

            try:

                # --------------------------------------------
                # MOBILE
                # --------------------------------------------

                if mobile_index >= len(row):

                    raise ValueError(
                        "Mobile column is missing in row."
                    )

                mobile = normalize_mobile(
                    row[mobile_index]
                )

                if not mobile:

                    raise ValueError(
                        "Mobile number is empty."
                    )

                # --------------------------------------------
                # DUPLICATE
                # --------------------------------------------

                if (
                    SKIP_DUPLICATE_MOBILES
                    and mobile in processed_mobiles
                ):

                    logs.append(
                        SMSLog(
                            campaign=campaign,
                            mobile=mobile,
                            message="SKIPPED",
                            status="skipped",
                            response=(
                                "Duplicate mobile "
                                "number in Excel."
                            ),
                        )
                    )

                    skipped += 1

                    continue

                if SKIP_DUPLICATE_MOBILES:

                    processed_mobiles.add(
                        mobile
                    )

                # --------------------------------------------
                # MESSAGE
                # --------------------------------------------

                message = build_message(
                    template.content,
                    variables,
                    variable_indexes,
                    row,
                )

                # --------------------------------------------
                # SEND
                # --------------------------------------------

                success, response = send_sms(
                    session,
                    mobile,
                    message,
                    template.template_id,
                )

                # --------------------------------------------
                # STATUS
                # --------------------------------------------

                if success:

                    status = "sent"
                    sent += 1

                else:

                    status = "failed"
                    failed += 1

                # --------------------------------------------
                # LOG
                # --------------------------------------------

                logs.append(
                    SMSLog(
                        campaign=campaign,
                        mobile=mobile,
                        message=message,
                        status=status,
                        response=response,
                    )
                )

            except Exception as exc:

                failed += 1

                mobile_value = ""

                if (
                    mobile_index < len(row)
                    and row[mobile_index] is not None
                ):

                    mobile_value = normalize_mobile(
                        row[mobile_index]
                    )

                logs.append(
                    SMSLog(
                        campaign=campaign,
                        mobile=mobile_value,
                        message="ERROR",
                        status="failed",
                        response=str(exc),
                    )
                )

            # =================================================
            # BULK LOG INSERT
            # =================================================

            if len(logs) >= LOG_BATCH_SIZE:

                logs = flush_logs(logs)

            # =================================================
            # PROGRESS
            # =================================================

            if (
                processed
                % PROGRESS_BATCH_SIZE
                == 0
            ):

                update_campaign_progress(
                    campaign_id,
                    processed,
                    sent,
                    failed,
                    skipped,
                )

        # ====================================================
        # FINAL LOG INSERT
        # ====================================================

        if logs:

            logs = flush_logs(logs)

        # ====================================================
        # COMPLETED
        # ====================================================

        SMSCampaign.objects.filter(
            id=campaign_id,
        ).update(
            processed=processed,
            sent=sent,
            failed=failed,
            skipped=skipped,
            status="COMPLETED",
            error_message=None,
            completed_at=timezone.now(),
        )

        return {
            "status": "COMPLETED",
            "campaign_id": campaign_id,
            "total": processed,
            "sent": sent,
            "failed": failed,
            "skipped": skipped,
        }

    except Exception as exc:

        # ----------------------------------------------
        # Save remaining logs
        # ----------------------------------------------

        try:

            if logs:
                flush_logs(logs)

        except Exception:
            pass

        # ----------------------------------------------
        # Campaign failure
        # ----------------------------------------------

        fail_campaign(
            campaign_id,
            exc,
        )

        raise

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass

        if excel_file is not None:

            try:
                excel_file.close()
            except Exception:
                pass

        if session is not None:

            try:
                session.close()
            except Exception:
                pass